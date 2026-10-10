"""Per-campaign reports: a summary, one row per recipient, and link clicks.

Downloadable as an Excel workbook (three sheets) or as CSV (one part per file).
All times are UTC.
"""
import io
import re
from collections import defaultdict

from app import db
from app.models import EmailEvent, EmailLog, Subscriber
from app.services import campaign_service
from app.utils.helpers import ServiceError, parse_datetime, rows_to_csv, utcnow

FORMATS = ('xlsx', 'csv')
PARTS = ('recipients', 'summary', 'links')

RECIPIENT_COLUMNS = [
    ('email', 'Email'), ('first_name', 'First name'), ('last_name', 'Last name'),
    ('delivery_status', 'Delivery status'), ('subject_line', 'Subject'), ('variant', 'A/B variant'),
    ('sent_at', 'Sent at (UTC)'), ('opened', 'Opened'), ('opens', 'Opens'), ('first_opened_at', 'First opened (UTC)'),
    ('clicked', 'Clicked'), ('clicks', 'Clicks'), ('first_clicked_at', 'First clicked (UTC)'),
    ('links_clicked', 'Links clicked'), ('bounced', 'Bounce'), ('unsubscribed', 'Unsubscribed'),
    ('error', 'Error'), ('subscriber_status', 'Subscriber status now'),
]
LINK_COLUMNS = [('url', 'Link'), ('clicks', 'Clicks'), ('unique_clicks', 'Unique clickers'),
                ('click_share', 'Share of clicks')]


def build_report(business_id, campaign_id):
    analytics = campaign_service.get_campaign_analytics(business_id, campaign_id)
    campaign = analytics['campaign']

    clicked_links = defaultdict(list)
    link_totals = defaultdict(lambda: [0, set()])
    unsubscribed = set()
    events = db.session.query(EmailEvent.email_log_id, EmailEvent.event_type, EmailEvent.url).filter(
        EmailEvent.business_id == business_id, EmailEvent.campaign_id == campaign_id,
        EmailEvent.event_type.in_(['click', 'unsubscribe']))
    for log_id, event_type, url in events:
        if event_type == 'unsubscribe':
            unsubscribed.add(log_id)
        elif url:
            if url not in clicked_links[log_id]:
                clicked_links[log_id].append(url)
            link_totals[url][0] += 1
            link_totals[url][1].add(log_id)

    recipients = []
    query = db.session.query(EmailLog, Subscriber).outerjoin(Subscriber, Subscriber.id == EmailLog.subscriber_id) \
        .filter(EmailLog.business_id == business_id, EmailLog.campaign_id == campaign_id) \
        .order_by(EmailLog.recipient_email)
    for log, subscriber in query.yield_per(1000):
        recipients.append({
            'email': log.recipient_email,
            'first_name': subscriber.first_name if subscriber else None,
            'last_name': subscriber.last_name if subscriber else None,
            'delivery_status': log.status,
            'subject_line': log.subject_line,
            'variant': chr(ord('A') + log.subject_variant) if log.subject_variant is not None else None,
            'sent_at': log.sent_at,
            'opened': 'yes' if log.opened_at else 'no',
            'opens': log.open_count or 0,
            'first_opened_at': log.opened_at,
            'clicked': 'yes' if log.clicked_at else 'no',
            'clicks': log.click_count or 0,
            'first_clicked_at': log.clicked_at,
            'links_clicked': ' | '.join(clicked_links.get(log.id, [])) or None,
            'bounced': log.bounce_type or ('yes' if log.bounced_at else None),
            'unsubscribed': 'yes' if log.id in unsubscribed else 'no',
            'error': log.error_message,
            'subscriber_status': subscriber.status if subscriber else 'deleted',
        })

    total_link_clicks = sum(total for total, _ in link_totals.values())
    links = sorted(({'url': url, 'clicks': total, 'unique_clicks': len(logs),
                     'click_share': total / total_link_clicks if total_link_clicks else 0}
                    for url, (total, logs) in link_totals.items()), key=lambda r: -r['clicks'])

    rate = lambda key: analytics[key] / 100  # noqa: E731  (stored as percentages)
    summary = [
        ('Campaign', campaign['name'], None),
        ('Subject', campaign.get('subject_line') or '(template subject)', None),
        ('Status', campaign['status'], None),
        ('Audience', campaign.get('audience_summary'), None),
        ('Sent at (UTC)', _parse(campaign.get('send_time')), 'date'),
        ('Recipients', analytics['recipients'] or 0, 'int'),
        ('Sent', analytics['sent'], 'int'),
        ('Still queued', analytics['pending'], 'int'),
        ('Failed', analytics['failed'], 'int'),
        ('Unique opens', analytics['unique_opens'], 'int'),
        ('Open rate', rate('open_rate'), 'pct'),
        ('Total opens', analytics['total_opens'], 'int'),
        ('Unique clicks', analytics['unique_clicks'], 'int'),
        ('Click rate', rate('click_rate'), 'pct'),
        ('Click-to-open rate', rate('click_to_open_rate'), 'pct'),
        ('Total clicks', analytics['total_clicks'], 'int'),
        ('Bounces', analytics['bounces'] or 0, 'int'),
        ('Bounce rate', rate('bounce_rate'), 'pct'),
        ('Unsubscribes', analytics['unsubscribes'] or 0, 'int'),
        ('Unsubscribe rate', rate('unsubscribe_rate'), 'pct'),
    ]
    for variant in (analytics.get('ab_test') or {}).get('variants', []):
        letter = chr(ord('A') + variant['variant'])
        summary += [
            (f'Variant {letter} subject', variant['subject_line'], None),
            (f'Variant {letter} sent', variant['sent'], 'int'),
            (f'Variant {letter} open rate', variant['open_rate'] / 100, 'pct'),
            (f'Variant {letter} click rate', variant['click_rate'] / 100, 'pct'),
        ]
    summary.append(('Report generated (UTC)', utcnow().replace(microsecond=0), 'date'))
    return {'campaign': campaign, 'summary': summary, 'recipients': recipients, 'links': links}


def _parse(value):
    return parse_datetime(value) if value else None


def _text(value):
    if value is None:
        return ''
    if hasattr(value, 'strftime'):
        return value.strftime('%Y-%m-%d %H:%M:%S')
    return value


def filename(report, ext, part=None):
    slug = re.sub(r'[^A-Za-z0-9]+', '-', report['campaign']['name']).strip('-').lower()[:60] or 'campaign'
    return f"{slug}-{part + '-' if part else ''}report-{utcnow():%Y%m%d}.{ext}"


def to_csv(report, part='recipients'):
    if part not in PARTS:
        raise ServiceError(f'part must be one of: {", ".join(PARTS)}')
    if part == 'summary':
        rows = [{'Metric': label, 'Value': (f'{value * 100:.2f}%' if kind == 'pct' else _text(value))}
                for label, value, kind in report['summary']]
        return rows_to_csv(rows, ['Metric', 'Value'])
    columns = RECIPIENT_COLUMNS if part == 'recipients' else LINK_COLUMNS
    rows = []
    for row in report[part]:
        out = {}
        for key, label in columns:
            value = row.get(key)
            out[label] = f'{value * 100:.2f}%' if key == 'click_share' else _text(value)
        rows.append(out)
    return rows_to_csv(rows, [label for _, label in columns])


def to_xlsx(report):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='047857')
    date_format = 'yyyy-mm-dd hh:mm'

    def append(sheet, values):
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            # Always text: a value starting with "=" from subscriber data must not become a formula
            if isinstance(cell.value, str) and cell.value.startswith('='):
                cell.data_type = 's'

    def header(sheet, labels, widths):
        sheet.append(labels)
        for index, cell in enumerate(sheet[1], start=1):
            cell.font, cell.fill = header_font, header_fill
            cell.alignment = Alignment(vertical='center')
            sheet.column_dimensions[cell.column_letter].width = widths[index - 1]
        sheet.freeze_panes = 'A2'

    workbook = Workbook()
    summary = workbook.active
    summary.title = 'Summary'
    header(summary, ['Metric', 'Value'], [28, 60])
    for label, value, kind in report['summary']:
        append(summary, [label, value])
        cell = summary.cell(row=summary.max_row, column=2)
        cell.alignment = Alignment(horizontal='left')
        if kind == 'pct':
            cell.number_format = '0.00%'
        elif kind == 'date':
            cell.number_format = date_format

    sheet = workbook.create_sheet('Recipients')
    widths = {'email': 32, 'subject_line': 36, 'links_clicked': 48, 'error': 40}
    header(sheet, [label for _, label in RECIPIENT_COLUMNS],
           [widths.get(key, 18 if key.endswith('_at') else 14) for key, _ in RECIPIENT_COLUMNS])
    date_columns = [i for i, (key, _) in enumerate(RECIPIENT_COLUMNS, start=1) if key.endswith('_at')]
    for row in report['recipients']:
        append(sheet, [row.get(key) for key, _ in RECIPIENT_COLUMNS])
        for column in date_columns:
            sheet.cell(row=sheet.max_row, column=column).number_format = date_format
    if report['recipients']:
        sheet.auto_filter.ref = sheet.dimensions

    sheet = workbook.create_sheet('Links')
    header(sheet, [label for _, label in LINK_COLUMNS], [70, 10, 16, 16])
    for row in report['links']:
        append(sheet, [row.get(key) for key, _ in LINK_COLUMNS])
        sheet.cell(row=sheet.max_row, column=4).number_format = '0.0%'

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
