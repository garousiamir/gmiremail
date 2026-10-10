import csv
import io

from openpyxl import load_workbook

from app import db
from app.models import Business, Campaign, EmailLog, Subscriber
from app.services import email_service, tracking_service
from tests.conftest import add_subscriber, create_template


def _sent_campaign(client, auth):
    add_subscriber(client, auth, 'ann@example.com', first_name='Ann')
    add_subscriber(client, auth, 'bob@example.com', first_name='=HYPERLINK("http://evil")')
    template = create_template(client, auth)
    campaign = client.post('/api/campaigns', headers=auth, json={
        'name': 'October news!', 'template_id': template['id']}).get_json()
    client.post(f"/api/campaigns/{campaign['id']}/send", json={}, headers=auth)
    email_service.process_email_queue()
    ann = EmailLog.query.filter_by(recipient_email='ann@example.com').one()
    tracking_service.track_email_open(ann.id)
    tracking_service.track_email_click(ann.id, 'https://example.com/a')
    tracking_service.track_email_click(ann.id, 'https://example.com/a')
    tracking_service.track_email_click(ann.id, 'https://example.com/b')
    return campaign


def _export(client, auth, campaign_id, **body):
    return client.post(f'/api/campaigns/{campaign_id}/export', json=body, headers=auth)


def test_campaign_report_xlsx(client, auth, outbox):
    campaign = _sent_campaign(client, auth)
    response = _export(client, auth, campaign['id'], format='xlsx')
    assert response.status_code == 200
    assert 'october-news-report-' in response.headers['Content-Disposition']
    book = load_workbook(io.BytesIO(response.data))
    assert book.sheetnames == ['Summary', 'Recipients', 'Links']

    summary = {row[0]: row[1] for row in book['Summary'].iter_rows(min_row=2, values_only=True)}
    assert summary['Campaign'] == 'October news!'
    assert summary['Sent'] == 2 and summary['Unique opens'] == 1 and summary['Unique clicks'] == 1
    assert summary['Open rate'] == 0.5 and summary['Total clicks'] == 3

    rows = list(book['Recipients'].iter_rows(values_only=True))
    header = rows[0]
    by_email = {r[0]: dict(zip(header, r)) for r in rows[1:]}
    ann = by_email['ann@example.com']
    assert ann['Opened'] == 'yes' and ann['Clicks'] == 3
    assert ann['Links clicked'] == 'https://example.com/a | https://example.com/b'
    assert by_email['bob@example.com']['Clicked'] == 'no'
    # Subscriber data that looks like a formula stays plain text
    cell = next(c for c in book['Recipients']['B'] if c.value and str(c.value).startswith('=HYPERLINK'))
    assert cell.data_type == 's'

    links = list(book['Links'].iter_rows(min_row=2, values_only=True))
    assert links[0][:3] == ('https://example.com/a', 2, 1)


def test_campaign_report_csv_parts(client, auth, outbox):
    campaign = _sent_campaign(client, auth)
    recipients = _export(client, auth, campaign['id'], format='csv', part='recipients')
    assert recipients.status_code == 200 and recipients.mimetype == 'text/csv'
    rows = list(csv.DictReader(io.StringIO(recipients.data.decode('utf-8-sig'))))
    assert {r['Email'] for r in rows} == {'ann@example.com', 'bob@example.com'}
    assert next(r for r in rows if r['Email'] == 'bob@example.com')['First name'].startswith("'=")

    summary = _export(client, auth, campaign['id'], format='csv', part='summary').data.decode('utf-8-sig')
    assert 'Open rate,50.00%' in summary
    links = _export(client, auth, campaign['id'], format='csv', part='links').data.decode('utf-8-sig')
    assert 'https://example.com/a,2,1,66.67%' in links

    assert _export(client, auth, campaign['id'], format='pdf').status_code == 400
    assert _export(client, auth, campaign['id'], format='csv', part='nope').status_code == 400


def test_campaign_report_is_private(client, auth, other_auth, outbox):
    campaign = _sent_campaign(client, auth)
    assert _export(client, other_auth, campaign['id'], format='xlsx').status_code == 404


def test_delete_account_removes_everything(client, auth, other_auth, outbox):
    _sent_campaign(client, auth)
    add_subscriber(client, other_auth, 'keep@example.com')
    me = client.get('/api/businesses/me', headers=auth).get_json()

    assert client.delete('/api/businesses/me', json={'password': 'wrong', 'confirm': 'DELETE'},
                         headers=auth).status_code == 403
    assert client.delete('/api/businesses/me', json={'password': 'supersecret', 'confirm': 'yes'},
                         headers=auth).status_code == 400
    assert db.session.get(Business, me['id']) is not None

    response = client.delete('/api/businesses/me', json={'password': 'supersecret', 'confirm': 'delete'}, headers=auth)
    assert response.status_code == 200
    db.session.expire_all()
    assert db.session.get(Business, me['id']) is None
    assert Subscriber.query.filter_by(business_id=me['id']).count() == 0
    assert Campaign.query.filter_by(business_id=me['id']).count() == 0
    assert EmailLog.query.filter_by(business_id=me['id']).count() == 0
    # The old token stops working; other accounts are untouched
    assert client.get('/api/businesses/me', headers=auth).status_code == 401
    assert Subscriber.query.filter_by(email='keep@example.com').count() == 1
    login = client.post('/api/auth/login', json={'email': 'owner@acme.example.com', 'password': 'supersecret'})
    assert login.status_code == 401
